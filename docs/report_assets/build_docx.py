from __future__ import annotations

from pathlib import Path
import re
import sys

sys.path.insert(0, "/tmp/safedrive-report-deps")

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


DOCS = Path(__file__).resolve().parent.parent
SOURCE = DOCS / "SAFEDRIVE_AI_FINAL_REPORT_AR.md"
DOCX = DOCS / "SAFEDRIVE_AI_FINAL_REPORT_AR.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = align
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        rtl = OxmlElement("w:rtl")
        run._element.get_or_add_rPr().append(rtl)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1); section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)
    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.35
    for style_name, size, color in [
        ("Title", 24, "17365D"), ("Heading 1", 19, "17365D"),
        ("Heading 2", 15, "24527A"), ("Heading 3", 13, "355C7D")
    ]:
        style = document.styles[style_name]
        style.font.name = "Arial"; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    if "CodeBlock" not in [s.name for s in document.styles]:
        style = document.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "DejaVu Sans Mono"; style.font.size = Pt(8.5)
        style.paragraph_format.left_indent = Cm(.5)
        style.paragraph_format.right_indent = Cm(.5)


def add_table(document: Document, lines: list[str]) -> None:
    rows = [[clean_inline(c) for c in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) >= 2 and all(set(c.strip()) <= set(":-") for c in rows[1]):
        rows.pop(1)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            if i == 0:
                set_cell_shading(cell, "D9EAF7")
            for p in cell.paragraphs:
                set_rtl(p)
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if i == 0: run.bold = True


def add_picture(document: Document, alt: str, path_text: str) -> None:
    path = (DOCS / path_text).resolve()
    if not path.is_file():
        p = document.add_paragraph(f"[صورة غير موجودة: {path_text}]")
        set_rtl(p)
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.3))
    caption = document.add_paragraph(clean_inline(alt))
    set_rtl(caption, WD_ALIGN_PARAGRAPH.CENTER)
    for run in caption.runs:
        run.italic = True; run.font.size = Pt(9.5); run.font.color.rgb = RGBColor(80,80,80)


def add_toc(document: Document) -> None:
    p = document.add_paragraph()
    set_rtl(p)
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    msg = OxmlElement("w:t"); msg.text = "حدّث الفهرس في Word بالضغط على F9"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, msg, end])


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document(); configure(document)
    in_frontmatter = lines and lines[0].strip() == "---"
    frontmatter_done = False
    in_code = False; code_lines: list[str] = []
    i = 0; h1_count = 0
    while i < len(lines):
        raw = lines[i]; line = raw.rstrip()
        if in_frontmatter:
            if line == "---" and i > 0:
                in_frontmatter = False; frontmatter_done = True
            i += 1; continue
        if line.startswith("<div") or line.startswith("</div") or line == "<br>":
            i += 1; continue
        if line.startswith("```"):
            if in_code:
                p = document.add_paragraph("\n".join(code_lines), style="CodeBlock")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                in_code = False; code_lines = []
            else:
                in_code = True; code_lines = []
            i += 1; continue
        if in_code:
            code_lines.append(line); i += 1; continue
        if line.startswith("|") and line.endswith("|"):
            table_lines=[]
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i]); i += 1
            add_table(document, table_lines); continue
        mimg = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if mimg:
            add_picture(document, mimg.group(1), mimg.group(2)); i += 1; continue
        mh = re.match(r"^(#{1,3})\s+(.*)$", line)
        if mh:
            level=len(mh.group(1)); title=clean_inline(mh.group(2))
            if level == 1:
                if h1_count > 0: document.add_page_break()
                h1_count += 1
            p=document.add_paragraph(title, style=f"Heading {level}")
            set_rtl(p)
            if title == "فهرس المحتويات": add_toc(document)
            i += 1; continue
        if line.strip() in {"---", ""}:
            i += 1; continue
        if re.match(r"^[-*]\s+", line):
            p=document.add_paragraph(clean_inline(re.sub(r"^[-*]\s+", "", line)), style="List Bullet")
            set_rtl(p); i += 1; continue
        if re.match(r"^\d+\.\s+", line):
            p=document.add_paragraph(clean_inline(re.sub(r"^\d+\.\s+", "", line)), style="List Number")
            set_rtl(p); i += 1; continue
        if line.startswith(">"):
            p=document.add_paragraph(clean_inline(line.lstrip("> ")))
            p.paragraph_format.right_indent=Cm(.5); set_rtl(p)
            i += 1; continue
        # Join wrapped prose lines until the next structural line.
        parts=[line.strip()]; j=i+1
        while j < len(lines):
            nxt=lines[j].rstrip()
            if (not nxt or nxt.startswith(("#","```","|","!","<",">","---"))
                    or re.match(r"^[-*]\s+",nxt) or re.match(r"^\d+\.\s+",nxt)):
                break
            parts.append(nxt.strip()); j += 1
        p=document.add_paragraph(clean_inline(" ".join(parts))); set_rtl(p)
        i=j

    props = document.core_properties
    props.title = "SafeDrive AI: A Real-Time Driver Monitoring and Risky Behavior Detection System"
    props.author = "فريق SafeDrive AI — جامعة دمشق"
    props.subject = "تقرير مشروع تخرج — نظام مراقبة السائق"
    document.save(DOCX)
    if DOCX.stat().st_size < 100_000:
        raise RuntimeError("DOCX is unexpectedly small")
    print(f"created {DOCX} ({DOCX.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
